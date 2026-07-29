#!/usr/bin/env python3
"""AI电商行业研究报告 — 正式版PPT生成脚本（16:9）。"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
CHARTS = ROOT / "charts"
OUT = ROOT / "report"
OUT.mkdir(exist_ok=True)

NAVY = RGBColor(0x0F, 0x27, 0x44)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
INK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0E, 0x74, 0x90)

FONT = "Microsoft YaHei"
SW, SH = Inches(13.333), Inches(7.5)
outline = []


def set_font(run, size=14, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def box(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for text, size, bold, color in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        run = p.add_run()
        run.text = text
        set_font(run, size=size, bold=bold, color=color)
        p.space_after = Pt(max(3, size * 0.3))
    return tb


def rect(slide, x, y, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def slide_base(prs, chapter, title, page):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, SW, Inches(0.08), BLUE)
    box(slide, Inches(0.7), Inches(0.28), Inches(12), Inches(0.35),
        [(chapter, 11, True, ACCENT)])
    box(slide, Inches(0.7), Inches(0.58), Inches(12), Inches(0.7),
        [(title, 20, True, NAVY)])
    box(slide, Inches(0.7), Inches(7.1), Inches(12), Inches(0.3),
        [(f"AI电商行业研究报告｜正式版 v2.0｜数据截至2026-07｜{page}", 8.5, False, GRAY)])
    outline.append(f"\n## 第{page}页｜{chapter}｜{title}")
    return slide


def bullets(slide, items, y=Inches(1.45), size=13.5):
    lines = []
    for it in items:
        if isinstance(it, tuple):
            lines.append((f"● {it[0]}", size, True, INK))
            if it[1]:
                lines.append((f"   {it[1]}", size - 1.5, False, GRAY))
        else:
            lines.append((f"● {it}", size, False, INK))
    box(slide, Inches(0.75), y, Inches(11.9), Inches(5.3), lines)
    for it in items:
        outline.append(f"  - {it[0]}：{it[1]}" if isinstance(it, tuple) else f"  - {it}")


def note(slide, text):
    box(slide, Inches(0.75), Inches(6.55), Inches(11.9), Inches(0.45),
        [(text, 9, False, GRAY)])
    outline.append(f"  注：{text}")


def add_image(slide, img, y=Inches(1.4), max_h=4.9):
    from PIL import Image
    path = CHARTS / img
    with Image.open(path) as im:
        w_px, h_px = im.size
    disp_h = Inches(max_h)
    disp_w = Emu(int(disp_h * w_px / h_px))
    max_w = Inches(12.0)
    if disp_w > max_w:
        disp_w = max_w
        disp_h = Emu(int(disp_w * h_px / w_px))
    left = Emu(int((SW - disp_w) / 2))
    slide.shapes.add_picture(str(path), left, y, width=disp_w, height=disp_h)
    outline.append(f"  [图：charts/{img}]")


def table_slide(prs, chapter, title, page, headers, rows, widths, source, fs=10.5):
    slide = slide_base(prs, chapter, title, page)
    n_r, n_c = len(rows) + 1, len(headers)
    tbl = slide.shapes.add_table(n_r, n_c, Inches(0.7), Inches(1.45),
                                 Inches(sum(widths)), Inches(0.48 * n_r)).table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = h
        set_font(run, size=fs, bold=True, color=WHITE)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if i % 2 else WHITE
            cell.margin_top = Pt(3)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            set_font(run, size=fs - (1 if len(str(val)) > 55 else 0), color=INK)
    note(slide, source)
    outline.append("  [表] " + " | ".join(headers))
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    p = 0

    # 1 封面
    p += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, NAVY)
    rect(s, Inches(0.9), Inches(2.1), Inches(0.12), Inches(1.8), ACCENT)
    box(s, Inches(1.25), Inches(2.0), Inches(11), Inches(2.2), [
        ("AI电商行业研究报告", 36, True, WHITE),
        ("全球与中国市场全景 · 双路径分析 · 产业链与场景机会", 16, False, RGBColor(0xBF, 0xDB, 0xFE)),
    ])
    box(s, Inches(1.25), Inches(4.6), Inches(11), Inches(1.8), [
        ("正式版 v2.0", 14, True, WHITE),
        ("数据截至：2026年7月", 12, False, RGBColor(0x93, 0xC5, 0xFD)),
        ("数据原则：仅采用第三方监测、公司财报或官方公告可追溯数据；未经独立验证的自报数据不予展示", 11, False, RGBColor(0x93, 0xC5, 0xFD)),
    ])
    outline.append("## 第1页｜封面｜AI电商行业研究报告（正式版v2.0）")

    # 2 目录
    p += 1
    s = slide_base(prs, "目录", "报告结构", p)
    bullets(s, [
        "第一章  研究说明与数据口径",
        "第二章  执行摘要：主要发现与建议方向",
        "第三章  全球市场全景：规模、导流与假日季验证",
        "第四章  双路径分析：站外LLM与站内AI导购",
        "第五章  中国市场格局",
        "第六章  产业链结构（MECE）",
        "第七章  零售与生鲜场景结合点",
        "第八章  产品技术机会与结论建议",
    ], size=15)

    # 3 研究说明
    p += 1
    s = slide_base(prs, "第一章  研究说明", "研究范围、能力分级与数据采用原则", p)
    bullets(s, [
        ("研究范围", "AI介入消费者需求唤起—决策—交易—履约—售后链路，或介入商家经营链路的电商形态"),
        ("路径A：站外2C LLM", "通用AI助手内完成购物决策起点：ChatGPT、Gemini、Perplexity、豆包、千问、元宝"),
        ("路径B：站内AI导购", "电商/零售App内嵌助手：Amazon Rufus、Walmart Sparky、淘宝AI万能搜、Instacart Cart Assistant"),
        ("能力分级A1—A5", "站内问答 → 站外种草导流 → 对话内闭环 → 授权代理执行 → 自主代理"),
        ("数据原则", "优先第三方与财报数据；公司自报须来自财报电话会或官方公告；未经独立验证的自报数据不展示；严格区分「AI影响的销售」与「AI平台内成交」"),
    ], size=13)

    # 4 执行摘要
    p += 1
    s = slide_base(prs, "第二章  执行摘要", "三项主要发现", p)
    bullets(s, [
        ("发现一：决策入口向对话界面迁移",
         "美国零售AI导流2026Q1同比+393%，转化率一年内由-38%转为+42%（Adobe）；假日季AI影响约20%全球零售订单、2620亿美元（Salesforce）；国内豆包MAU 3.83亿、千问1.66亿（QuestMobile）"),
        ("发现二：交易宜留在零售商自有体系",
         "ChatGPT Instant Checkout上线约半年后收缩；Sparky以插件接入外部AI、结账回自有体系后，转化约为自有站70%，高于平台代结账阶段（约1/3）"),
        ("发现三：站内导购效果可量化，生鲜与会员制适配度更高",
         "Rufus 2025年使用用户超3亿，购买完成率+60%，约120亿美元增量年化销售（亚马逊财报电话会，公司口径）"),
    ], size=13)
    note(s, "数据来源详见各章节图表脚注；本页仅列已纳入正式报告的可追溯数据。")

    # 5 建议
    p += 1
    s = slide_base(prs, "第二章  执行摘要", "三项建议方向", p)
    bullets(s, [
        ("站内先行", "在自有App/小程序落地高确定性场景：商品问答、菜谱/清单成车、售后自动化；以转化率与使用率为考核指标"),
        ("供给侧就绪", "推进商品数据结构化与大模型可读性改造；对主要AI入口建立「被推荐率/被平替率」月度监测"),
        ("站外合作坚持交易主权", "与通用AI入口合作时采用「插件/小程序跳转、结账留在自有体系」；避免将结账权让渡给第三方AI平台"),
    ], size=14)

    # 6-9 全球数据图
    p += 1
    s = slide_base(prs, "第三章  全球市场 1/4", "图1  代理式商务规模预测：口径差异决定结论差异", p)
    add_image(s, "01_global_agentic_forecasts.png", max_h=5.0)

    p += 1
    s = slide_base(prs, "第三章  全球市场 2/4", "图2  美国零售网站AI导流增长", p)
    add_image(s, "02_us_ai_traffic_growth.png", max_h=5.0)

    p += 1
    s = slide_base(prs, "第三章  全球市场 3/4", "图3  AI导流质量：转化率与单次访问收入逆转", p)
    add_image(s, "03_ai_conversion_reversal.png", max_h=5.0)

    p += 1
    s = slide_base(prs, "第三章  全球市场 4/4", "图6  2025假日季：AI影响约20%全球零售订单", p)
    add_image(s, "06_holiday_ai_influence.png", max_h=5.0)

    # 10 关键案例
    p += 1
    s = slide_base(prs, "第三章  全球市场", "行业案例：对话内代结账受挫，插件+自有结账更稳妥", p)
    bullets(s, [
        ("Instant Checkout（2025-09至约2026-03）",
         "ChatGPT内完成结账的试点上线后收缩；公开报道显示转化明显低于商家自有渠道，错购与体验问题突出"),
        ("Sparky插件模式（2026-03起）",
         "沃尔玛自有助手以插件形式进入ChatGPT/Gemini：对话内选品，回自有体系结账；转化约为自有站70%"),
        ("Amazon Rufus",
         "2025年使用用户超3亿；购买完成率+60%；约120亿美元增量年化销售（亚马逊财报电话会，公司口径）"),
        ("Instacart Cart Assistant",
         "企业级白标AI套件输出给Kroger、Sprouts等，覆盖膳食规划与对话建车（Instacart 2025-11新闻稿）"),
    ], size=12.5)
    note(s, "来源：OpenAI公告、媒体交叉报道、亚马逊财报电话会、Instacart官方新闻稿。")

    # 11 双路径
    p += 1
    table_slide(prs, "第四章  双路径分析", "路径A与路径B对照", p,
                ["维度", "路径A：站外2C LLM", "路径B：站内AI导购"],
                [
                    ["核心价值", "截获购物决策起点，影响品类与品牌心智", "降低站内决策成本，提升转化与客单"],
                    ["代表产品", "ChatGPT、Gemini、豆包、千问、元宝", "Rufus、Sparky、淘宝AI万能搜、Cart Assistant"],
                    ["变现方式", "交易佣金、归因结算、联盟分佣、订阅", "不直接收费；体现为GMV增量与留存"],
                    ["主要风险", "推荐公正性、幻觉错购、交易权争夺", "推理成本、答案页与广告位冲突"],
                    ["终局角色", "发现层与部分闭环入口", "交易场内默认决策层"],
                ],
                [1.8, 5.1, 5.1],
                "整理自公开产品进展与本报告第四章；量化数据见后续图表。", fs=11)

    # 12 入口规模
    p += 1
    s = slide_base(prs, "第四章  双路径分析", "图4  主要AI购物入口用户规模（可追溯数据）", p)
    add_image(s, "04_ai_entrance_user_scale.png", max_h=5.0)

    # 13 效果分口径
    p += 1
    s = slide_base(prs, "第四章  双路径分析", "图7  站内导购与站外引荐：分口径效果对照", p)
    add_image(s, "07_instore_ai_scorecard.png", max_h=5.0)

    # 14 中国宏观
    p += 1
    s = slide_base(prs, "第五章  中国市场", "图5  中国AI私域电商市场规模与渗透率", p)
    add_image(s, "05_china_ai_private_ecommerce.png", max_h=5.0)

    # 15 中国格局表
    p += 1
    table_slide(prs, "第五章  中国市场", "入口格局：四大阵营", p,
                ["阵营", "模式", "进展", "数据来源"],
                [
                    ["字节：豆包×抖音", "垂直闭环", "MAU 3.83亿；买前问豆包一级入口；订单纳入抖音电商归因", "QuestMobile；官方/媒体"],
                    ["阿里：千问×淘宝", "垂直闭环", "MAU 1.66亿；对话完成淘宝全链路", "QuestMobile；阿里官宣"],
                    ["腾讯×京东：元宝", "生态联盟", "对话出商品卡，跳转京东小程序成交", "联合官宣"],
                    ["内容与其他", "社区AI/站内工具", "小红书点点并入主站；拼多多、快手已上线AI搜索", "36氪、网经社等"],
                ],
                [2.2, 1.6, 4.8, 3.4],
                "仅列可追溯进展；未纳入未经独立验证的自报运营规模数据。", fs=10.5, )

    # 16 产业链
    p += 1
    table_slide(prs, "第六章  产业链", "六层结构（按交易功能环节切分）", p,
                ["层级", "职能", "收入模式", "主要成本"],
                [
                    ["L1 模型与算力", "推理与多模态能力", "API、云捆绑、订阅", "训练与推理算力"],
                    ["L2 入口与流量", "承接购物意图", "佣金、归因、订阅", "推理、获客、合规"],
                    ["L3 商品与交易", "商品池、交易与售后", "GMV与平台佣金", "运营、渠道费率、feed改造"],
                    ["L4 交易基建", "协议与支付信任", "通道费、令牌服务", "风控与标准建设"],
                    ["L5 商家服务", "GEO、SaaS、数字人、监测", "订阅与咨询", "研发与多模型适配"],
                    ["L6 履约供应链AI", "预测、补货、调度", "效率带来的成本节约", "数据与算法团队"],
                ],
                [2.3, 3.2, 3.3, 3.2],
                "竞合模式：垂直闭环 / 生态联盟与插件 / 中立第三方。详见报告第六章。", fs=11)

    # 17 场景
    p += 1
    s = slide_base(prs, "第七章  场景结合点", "综合零售旅程 × 生鲜结构性特征", p)
    box(s, Inches(0.7), Inches(1.4), Inches(6.0), Inches(5.0), [
        ("综合零售（按旅程）", 14, True, NAVY),
        ("需求唤起：站外LLM商品推荐", 12, False, INK),
        ("方案生成：答案报告式导购", 12, False, INK),
        ("选品比较：参数与评论提炼", 12, False, INK),
        ("比价凑单：算优惠、盯价", 12, False, INK),
        ("下单支付：闭环 vs 跳转自有结账", 12, False, INK),
        ("履约售后：物流问答与退换自动化", 12, False, INK),
        ("复购：偏好记忆与补货提醒", 12, False, INK),
    ])
    box(s, Inches(6.9), Inches(1.4), Inches(5.8), Inches(5.0), [
        ("生鲜零售（三类优先组合）", 14, True, NAVY),
        ("1. 餐桌助手", 12, True, INK),
        ("对话/语音 + 菜谱成车 + 周期购", 11.5, False, GRAY),
        ("2. 损耗联动导购", 12, True, INK),
        ("效期/库存接入推荐，临期与菜谱联动", 11.5, False, GRAY),
        ("3. 白标AI基建", 12, True, INK),
        ("面向区域商超输出对话建车能力", 11.5, False, GRAY),
        ("（参照Instacart Cart Assistant模式）", 11.5, False, GRAY),
    ])
    outline.append("  - 综合零售旅程七点；生鲜三类组合")

    # 18 产品机会
    p += 1
    s = slide_base(prs, "第八章  产品技术机会", "分角色机会与合规要点", p)
    bullets(s, [
        ("平台", "站内助手从问答升级到任务代理；AI答案页商业化规则；用户授权中心（额度/品类/撤销）"),
        ("品牌与商家", "商品feed结构化与大模型可读性；GEO监测与内容供给；自有助手接入主流AI入口"),
        ("生鲜与即时零售", "餐桌助手、损耗联动推荐、白标能力输出"),
        ("合规", "生成式AI与算法推荐备案、内容标识、个人信息保护影响评估、数据出境、未成年人与代扣授权、价格表述合规"),
    ], size=13)

    # 19 结论
    p += 1
    s = slide_base(prs, "第八章  结论与建议", "结论四点与行动优先级", p)
    bullets(s, [
        ("结论", "①规模化应用阶段已到；②站外LLM与站内导购长期并存；③「AI发现+零售商成交」更稳健；④生鲜与会员制适配度更高"),
        ("P0 监测与口径", "区分「影响」与「成交」；对主要AI入口做月度可见度监测"),
        ("P1 站内试点", "商品问答、清单/菜谱成车、售后自动化；设定转化与客诉门槛"),
        ("P2 供给就绪", "结构化feed与内容治理，同时服务站内助手与站外被发现"),
        ("P3 站外接入", "优先插件/小程序跳转、结账自有；明确数据与会员边界"),
        ("P4 生鲜专项", "评估餐桌助手与损耗联动推荐的投入产出"),
    ], size=12.5)

    # 20 附录
    p += 1
    table_slide(prs, "附录", "图表与数据来源索引", p,
                ["图号", "内容", "主要来源", "口径要点"],
                [
                    ["图1", "规模预测对比", "eMarketer/MS/Bain/McKinsey等", "窄=平台内结账；宽=AI影响/编排"],
                    ["图2", "美国AI导流增速", "Adobe Digital Insights 2026-04", "AI平台跳转至零售网站的访问"],
                    ["图3", "转化与RPV逆转", "Adobe Digital Insights", "AI相对非AI渠道"],
                    ["图4", "入口用户规模", "QuestMobile；OpenAI；Google；亚马逊财报", "WAU/MAU/年度用户不同，仅量级对照"],
                    ["图5", "中国AI私域电商", "网经社 2026-05", "私域子集，非全量AI电商GMV"],
                    ["图6", "假日季AI影响", "Salesforce 2025假日季报告", "「影响」宽口径"],
                    ["图7", "导购与引荐效果", "亚马逊财报；Salesforce；Adobe；Shopify", "两条证据线分列，不可合并"],
                ],
                [1.0, 2.4, 4.4, 4.2],
                "更新数据：修改 data/*.csv 后执行 python3 scripts/make_charts.py && python3 scripts/make_pptx.py",
                fs=10)

    out = OUT / "AI电商行业研究报告.pptx"
    prs.save(out)
    (OUT / "pptx_outline.md").write_text(
        "# PPT文字稿（自动生成）\n\n" + "\n".join(outline) + "\n", encoding="utf-8")
    print(f"已生成：{out.relative_to(ROOT)}（{p}页）")


if __name__ == "__main__":
    build()
